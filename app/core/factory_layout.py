"""Versioned presentation layouts. Scheduling events and the shared clock are immutable here."""
from __future__ import annotations
import copy
import io
import json
import math
import os
import threading
import time
import zipfile
from pathlib import Path
from xml.etree import ElementTree as ET

ROOT = Path(__file__).resolve().parents[2]
ASSETS = ROOT / 'data/factory_layout'
FIELDS = ['station_id','label','area','x','z','yaw','model_key','scale','show_model','label_height','animation','auto_dock','dock_x','dock_z','dock_yaw']
HEADERS = ['工位编号','显示名称','区域','X（米）','Z（米）','水平角度','模型编号','统一缩放','显示设备','标签高度','动画启用','自动停靠','停靠偏移X','停靠偏移Z','停靠朝向']

class LayoutError(ValueError):
    def __init__(self, errors):
        self.errors = errors
        super().__init__('；'.join(e['message'] for e in errors))

def issue(sid, field, message):
    return dict(station_id=sid, field=field, message=message)

def load_json(path):
    return json.loads(Path(path).read_text(encoding='utf-8-sig'))

def atomic_json(path, data):
    path = Path(path); path.parent.mkdir(parents=True, exist_ok=True)
    tmp = path.with_suffix('.tmp')
    tmp.write_text(json.dumps(data, ensure_ascii=False, indent=2, allow_nan=False), encoding='utf-8')
    os.replace(tmp, path)

def local_to_world(row, x, z):
    a = math.radians(row['yaw']); s = row['scale']
    return row['x'] + s*(x*math.cos(a)+z*math.sin(a)), row['z'] + s*(-x*math.sin(a)+z*math.cos(a))

def corners(row, default, models, pad=0):
    if row['model_key'] == default['model_key']:
        w,d = default['size_x'],default['size_z']
        cx,cz = default.get('center_x',0),default.get('center_z',0)
    else:
        model = models[row['model_key']]
        w,d = model['width_ratio']*default['base_width'], model['depth_ratio']*default['base_width']
        cx=cz=0
    w = w/2+pad/row['scale'];d=d/2+pad/row['scale']
    return [local_to_world(row,cx+x,cz+z) for x,z in [(-w,-d),(w,-d),(w,d),(-w,d)]]

def intersects(a,b):
    for polygon in (a,b):
        for p,q in zip(polygon,polygon[1:]+polygon[:1]):
            nx,nz = -(q[1]-p[1]),q[0]-p[0]
            aa=[x*nx+z*nz for x,z in a];bb=[x*nx+z*nz for x,z in b]
            if max(aa)<=min(bb)+1e-6 or max(bb)<=min(aa)+1e-6:return False
    return True

def box(x,z,w,d):return [(x-w/2,z-d/2),(x+w/2,z-d/2),(x+w/2,z+d/2),(x-w/2,z+d/2)]

def project(point,line):
    a,b=line;dx,dz=b[0]-a[0],b[1]-a[1]
    f=max(0,min(1,((point[0]-a[0])*dx+(point[1]-a[1])*dz)/(dx*dx+dz*dz)))
    return a[0]+f*dx,a[1]+f*dz

class LayoutStore:
    def __init__(self, directory=None, replay_path=None):
        self.directory=Path(directory or os.getenv('SIMPY_LAYOUT_STORE',ROOT/'results/factory_layout'))
        self.replay_path=Path(replay_path or os.getenv('SIMPY_REPLAY_PATH',ROOT/'results/unity_validation/replay.json'))
        self.lock=threading.RLock();self.clients={};self.route_cache={}

    @property
    def default(self):return load_json(ASSETS/'default.json')
    @property
    def current(self):
        path=self.directory/'current.json'
        return load_json(path) if path.exists() else self.default
    @property
    def replay(self):return load_json(self.replay_path)

    def parse(self, filename, content):
        if len(content)>2*1024*1024:raise LayoutError([issue('','file','布局文件不能超过 2 MiB')])
        if filename.lower().endswith('.json'):
            try:return json.loads(content.decode('utf-8-sig'))
            except (ValueError,UnicodeError) as e:raise LayoutError([issue('','file','JSON 格式错误：'+str(e))])
        if not filename.lower().endswith('.xlsx'):raise LayoutError([issue('','file','请选择 .xlsx 或 .json 文件')])
        try:
            with zipfile.ZipFile(io.BytesIO(content)) as z:
                if sum(i.file_size for i in z.infolist())>12*1024*1024:raise ValueError('解压后文件过大')
            from openpyxl import load_workbook
            book=load_workbook(io.BytesIO(content),read_only=True,data_only=False)
            sheet=book['工位布局'];rows=list(sheet.iter_rows(values_only=True))
            if not rows or list(rows[0][:len(HEADERS)])!=HEADERS:raise ValueError('请使用导出的原始表头')
            result=copy.deepcopy(self.current);result['stations']=[]
            for index,row in enumerate(rows[1:],2):
                if not any(v is not None for v in row):continue
                if any(isinstance(v,str) and v.startswith('=') for v in row):raise ValueError(f'第{index}行不能使用公式')
                result['stations'].append(dict(zip(FIELDS,row)))
            book.close();return result
        except Exception as e:raise LayoutError([issue('','file','Excel 读取失败：'+str(e))])

    def validate(self, document):
        default=self.default;current=self.current;models={m['key']:m for m in default['models']};base={r['station_id']:r for r in default['stations']}
        errors=[];normalized=copy.deepcopy(default);rows=[];seen=set()
        if not isinstance(document,dict):raise LayoutError([issue('','file','布局必须为 JSON 对象')])
        if document.get('schema_version')!=1 or document.get('catalog_version')!=default['catalog_version']:
            errors.append(issue('','version','布局格式或模型目录版本不匹配，请重新导出模板'))
        source=document.get('stations')
        if not isinstance(source,list):raise LayoutError([issue('','stations','缺少工位列表')])
        for raw in source:
            if not isinstance(raw,dict):errors.append(issue('','station_id','工位行必须为对象'));continue
            sid=str(raw.get('station_id','')).strip()
            if sid not in base or sid in seen:errors.append(issue(sid,'station_id','工位编号未知或重复：'+sid));continue
            seen.add(sid);row=copy.deepcopy(base[sid]);row.update({k:raw.get(k,row.get(k)) for k in FIELDS})
            for field in ['label','area','model_key']:
                value=row[field]
                if not isinstance(value,str) or not value.strip() or len(value)>100:errors.append(issue(sid,field,'文本不能为空且不能超过100字'))
                else:row[field]=value.strip()
            for field,lo,hi in [('x',-1000,1000),('z',-1000,1000),('yaw',-36000,36000),('scale',.25,3),('label_height',.1,5),('dock_x',-30,30),('dock_z',-30,30),('dock_yaw',-36000,36000)]:
                value=row[field]
                try:
                    if value is None and field.startswith('dock_'):value=0
                    if isinstance(value,bool):raise ValueError()
                    value=float(value)
                    if not math.isfinite(value) or not lo<=value<=hi:raise ValueError()
                    row[field]=value
                except (ValueError,TypeError):errors.append(issue(sid,field,f'{field} 必须是 {lo}～{hi} 范围内的有限数值'))
            for field in ['show_model','animation','auto_dock']:
                val=row[field]
                if val in (True,1,'是','true','TRUE'):row[field]=True
                elif val in (False,0,'否','false','FALSE'):row[field]=False
                else:errors.append(issue(sid,field,'开关请填写 是／否'))
            model=models.get(row['model_key']) if isinstance(row['model_key'],str) else None
            if not model:errors.append(issue(sid,'model_key','模型编号无效'))
            elif model.get('allowed_stations') and sid not in model['allowed_stations']:errors.append(issue(sid,'model_key','此模型仅支持 '+','.join(model['allowed_stations'])))
            rows.append(row)
        for sid in sorted(set(base)-seen):errors.append(issue(sid,'station_id','缺少工位：'+sid))
        if errors:raise LayoutError(errors)
        normalized['stations']=sorted(rows,key=lambda r:r['station_id']);normalized['revision']=current['revision']
        current_rows={r['station_id']:r for r in current['stations']}
        changes=[];changed_geometry=[];polygons={}
        geometry_fields=['x','z','yaw','scale','model_key']
        environment=load_json(ASSETS/'environment.json')
        for row in normalized['stations']:
            sid=row['station_id'];row['yaw']%=360;row['dock_yaw']%=360
            fields=[k for k in FIELDS if row[k]!=current_rows[sid].get(k)]
            if fields:changes.append(sid+'：'+', '.join(HEADERS[FIELDS.index(k)] for k in fields))
            poly=corners(row,base[sid],models);polygons[sid]=poly
            if any(row[k]!=base[sid][k] for k in geometry_fields):changed_geometry.append(sid)
            if any(not environment['min_x']<=x<=environment['max_x'] or not environment['min_z']<=z<=environment['max_z'] for x,z in poly):errors.append(issue(sid,'position','设备超出厂房可布置范围'))
        for sid in changed_geometry:
            poly=polygons[sid]
            for other,op in polygons.items():
                if other!=sid and intersects(poly,op):errors.append(issue(sid,'position','设备占地与 '+other+' 重叠'))
            for obstacle in environment['obstacles']+environment['roads']:
                if intersects(poly,obstacle['polygon']):errors.append(issue(sid,'position','设备占用 '+obstacle['name']))
        if errors:raise LayoutError(errors)
        snapshot=self.snapshot(normalized)
        return dict(valid=True,expected_revision=current['revision'],sim_id=snapshot['sim_id'],snapshot=snapshot,errors=[],changes=changes)

    def snapshot(self, layout=None):
        with self.lock:
            layout=layout or self.current;replay=self.replay
            key=(replay['sim_id'],json.dumps(layout['stations'],sort_keys=True))
            if key not in self.route_cache:
                self.route_cache={key:self.routes(layout,replay)}
            return dict(layout=layout,sim_id=replay['sim_id'],routes=self.route_cache[key])

    def routes(self,layout,replay):
        """Use the fixed road graph and conservative, obstacle-checked station approaches."""
        import heapq
        env=load_json(ASSETS/'environment.json');lines=env['lines']
        defaults={r['station_id']:r for r in self.default['stations']};models={m['key']:m for m in self.default['models']}
        blocked=[corners(r,defaults[r['station_id']],models,.68) for r in layout['stations']]+[o['polygon'] for o in env['obstacles']]
        buckets={};point_cache={}
        for poly in blocked:
            for x in range(math.floor(min(p[0] for p in poly)/5),math.floor(max(p[0] for p in poly)/5)+1):
                for z in range(math.floor(min(p[1] for p in poly)/5),math.floor(max(p[1] for p in poly)/5)+1):buckets.setdefault((x,z),[]).append(poly)
        def point_blocked(p):
            k=(round(p[0],5),round(p[1],5))
            if k not in point_cache:point_cache[k]=any(intersects(box(*p,.02,.02),poly) for poly in buckets.get((math.floor(p[0]/5),math.floor(p[1]/5)),[]))
            return point_cache[k]
        def segment_ok(a,b):
            steps=max(1,math.ceil(math.dist(a,b)/.3))
            return all(not point_blocked((a[0]+(b[0]-a[0])*i/steps,a[1]+(b[1]-a[1])*i/steps)) for i in range(steps+1))
        def approach(row):
            sid=row['station_id'];poly=corners(row,defaults[sid],models,.8)
            if row['auto_dock']:
                candidates=[]
                for a,b in zip(poly,poly[1:]+poly[:1]):
                    p=((a[0]+b[0])/2,(a[1]+b[1])/2)
                    if not point_blocked(p):candidates.append(p)
                if not candidates:raise LayoutError([issue(sid,'dock','工位周边没有安全停靠位置')])
            else:
                p=local_to_world(row,row['dock_x'],row['dock_z']);candidates=[p]
                if point_blocked(p):raise LayoutError([issue(sid,'dock','停靠点位于设备或固定障碍内')])
            options=[]
            for p in candidates:
                for line in lines:
                    q=project(p,line)
                    if segment_ok(p,q):options.append((math.dist(p,q),p,q))
            if not options:raise LayoutError([issue(sid,'dock','停靠点无法安全接入现有通道，请调整位置或停靠偏移')])
            _,p,q=min(options);return p,q
        docks={r['station_id']:approach(r) for r in layout['stations']}
        start=env['agv_start'];projections=[[tuple(a),tuple(b)] for a,b in lines]
        attachments={}
        for sid,(p,q) in docks.items():
            index=min(range(len(lines)),key=lambda i:math.dist(project(q,lines[i]),q));projections[index].append(q);attachments[sid]=(p,q)
        q=min((project(start,l) for l in lines),key=lambda q:math.dist(start,q));attachments['AGV_START']=(tuple(start),q)
        for i,l in enumerate(lines):
            if math.dist(project(q,l),q)<.01:projections[i].append(q)
            for j,m in enumerate(lines[:i]):
                if abs(l[0][1]-l[1][1])<.01 and abs(m[0][0]-m[1][0])<.01:p=(m[0][0],l[0][1])
                elif abs(m[0][1]-m[1][1])<.01 and abs(l[0][0]-l[1][0])<.01:p=(l[0][0],m[0][1])
                else:continue
                if math.dist(project(p,l),p)<1.5 and math.dist(project(p,m),p)<1.5:projections[i].append(p);projections[j].append(p)
        def key(p):return tuple(round(v,4) for v in p)
        def detour(a,b):
            if point_blocked(a) or point_blocked(b):return None
            if segment_ok(a,b):return [a,b]
            # A narrow corridor around the existing road, never a new factory-wide road network.
            xmin,xmax=min(a[0],b[0])-3.5,max(a[0],b[0])+3.5;zmin,zmax=min(a[1],b[1])-3.5,max(a[1],b[1])+3.5
            def grid(p):
                options=[(round(p[0]*2)/2+dx,round(p[1]*2)/2+dz) for dx in [-.5,0,.5] for dz in [-.5,0,.5]]
                return min((q for q in options if segment_ok(p,q)),key=lambda q:math.dist(q,p),default=None)
            start,end=grid(a),grid(b)
            if start is None or end is None:return None
            queue=[(math.dist(start,end),0,start)];costs={start:0};previous={}
            while queue:
                _,cost,u=heapq.heappop(queue)
                if u==end:break
                if cost>costs[u]:continue
                for dx,dz in [(0,.5),(0,-.5),(.5,0),(-.5,0)]:
                    v=(u[0]+dx,u[1]+dz)
                    if not xmin<=v[0]<=xmax or not zmin<=v[1]<=zmax or point_blocked(v):continue
                    if cost+.5<costs.get(v,float('inf')):costs[v]=cost+.5;previous[v]=u;heapq.heappush(queue,(cost+.5+math.dist(v,end),cost+.5,v))
            if end not in costs:return None
            path=[end]
            while path[-1]!=start:path.append(previous[path[-1]])
            path=[a]+list(reversed(path))+[b]
            simple=[path[0]]
            for i in range(1,len(path)-1):
                if not segment_ok(simple[-1],path[i+1]):simple.append(path[i])
            return simple+[b]
        graph={}
        for line,points in zip(lines,projections):
            points=sorted(set(map(key,points)),key=lambda p:math.dist(p,line[0]))
            for a,b in zip(points,points[1:]):
                path=detour(a,b)
                if path is None:continue
                length=sum(math.dist(u,v) for u,v in zip(path,path[1:]))
                graph.setdefault(a,[]).append((b,length,path));graph.setdefault(b,[]).append((a,length,list(reversed(path))))
        routes=[];cache={}
        for trip in replay['transfers']:
            source,target=trip['source_station'],trip['target_station'];pair=(source,target)
            if pair not in cache:
                p,a=attachments[source];r,b=attachments[target];a,b=key(a),key(b)
                queue=[(0,a)];dist={a:0};prev={}
                while queue:
                    cost,u=heapq.heappop(queue)
                    if u==b:break
                    if cost>dist[u]:continue
                    for v,c,path in graph.get(u,[]):
                        if cost+c<dist.get(v,float('inf')):dist[v]=cost+c;prev[v]=(u,path);heapq.heappush(queue,(cost+c,v))
                if b not in dist:raise LayoutError([issue(target,'dock','现有通道无法连接 '+source+' → '+target)])
                paths=[];node=b
                while node!=a:node,path=prev[node];paths.append(path)
                points=[p]+[point for path in reversed(paths) for point in path]+[r]
                cache[pair]=[dict(x=x,y=-.38,z=z) for x,z in points]
            routes.append(dict(trip_id=trip['trip_id'],route=cache[pair]))
        return routes

    def commit(self, document, expected_revision, sim_id):
        with self.lock:
            if self.current['revision']!=expected_revision or self.replay['sim_id']!=sim_id:raise RuntimeError('排产或布局已变化，请刷新预览')
            candidate=self.validate(document)['snapshot']['layout'];candidate['revision']=expected_revision+1
            if self.replay['sim_id']!=sim_id:raise RuntimeError('排产已变化，请刷新预览')
            previous=self.current
            self.directory.mkdir(parents=True,exist_ok=True)
            atomic_json(self.directory/'previous.json',previous)
            atomic_json(self.directory/f'revision_{candidate["revision"]}.json',candidate)
            atomic_json(self.directory/'current.json',candidate)
            return self.snapshot()

    def export(self, format):
        doc=self.current
        if format=='json':return json.dumps(doc,ensure_ascii=False,indent=2).encode('utf-8')
        # Preserve the authored Excel template and its validation/style; replace only input values.
        namespace='http://schemas.openxmlformats.org/spreadsheetml/2006/main';ET.register_namespace('',namespace)
        with zipfile.ZipFile(ASSETS/'template.xlsx') as source:
            root=ET.fromstring(source.read('xl/worksheets/sheet1.xml'))
            data=root.find('{'+namespace+'}sheetData')
            template_row=next(row for row in data if row.attrib['r']=='2')
            styles={cell.attrib['r'].rstrip('0123456789'):cell.attrib.get('s','0') for cell in template_row}
            for row in list(data):
                if int(row.attrib['r'])>1:data.remove(row)
            for index,row in enumerate(doc['stations'],2):
                element=ET.SubElement(data,'{'+namespace+'}row',r=str(index))
                for col,field in enumerate(FIELDS):
                    value=row[field];column=chr(65+col);cell=ET.SubElement(element,'{'+namespace+'}c',r=column+str(index),s=styles.get(column,'0'))
                    if isinstance(value,bool):value='是' if value else '否'
                    if isinstance(value,(int,float)):ET.SubElement(cell,'{'+namespace+'}v').text=str(value)
                    else:
                        cell.set('t','inlineStr');ET.SubElement(ET.SubElement(cell,'{'+namespace+'}is'),'{'+namespace+'}t').text=str(value)
            out=io.BytesIO()
            with zipfile.ZipFile(out,'w',zipfile.ZIP_DEFLATED) as dest:
                for item in source.infolist():dest.writestr(item,ET.tostring(root,encoding='utf-8',xml_declaration=True) if item.filename=='xl/worksheets/sheet1.xml' else source.read(item.filename))
        return out.getvalue()

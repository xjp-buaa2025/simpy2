"""Isolated physical placement store for the original 79-process Unity runtime."""
from __future__ import annotations

import copy
import hashlib
import io
import json
import math
import os
import threading
from pathlib import Path

from app.core.factory_layout import LayoutError, atomic_json, load_json, issue
from tools.pt6a29_integrated import ROOT, DEST, reroutes, world, retained_bounds

PROFILE_ID = 'pt6a29-original'
SLOT_FIELDS = ['id', 'zone_id', 'model_key', 'representative_id', 'x', 'z', 'yaw']
SLOT_HEADERS = ['设备实例编号', '所属工位', '原模型编号', '原模型来源工序', 'X（米）', 'Z（米）', '水平角度']


def metadata_equal(a, b):
    """Unity stores metres as float32; tolerate its roundtrip, not layout changes."""
    if isinstance(a, dict) and isinstance(b, dict):
        return a.keys() == b.keys() and all(metadata_equal(a[k], b[k]) for k in a)
    if isinstance(a, list) and isinstance(b, list):
        return len(a) == len(b) and all(metadata_equal(x, y) for x, y in zip(a, b))
    if isinstance(a, (int, float)) and not isinstance(a, bool) and isinstance(b, (int, float)) and not isinstance(b, bool):
        return math.isfinite(a) and math.isfinite(b) and abs(a-b) <= 0.0001
    return type(a) is type(b) and a == b


def signed_distance_to_zone(x, z, zone):
    left, top = world(zone['u'], zone['v'])
    right, bottom = world(zone['u'] + zone['width'], zone['v'] + zone['depth'])
    dx, dz = max(left - x, 0, x - right), max(bottom - z, 0, z - top)
    return math.hypot(dx, dz)


def route_errors(document):
    errors = []
    radius = document['agv_radius']
    auxiliaries=document.get('obstacles') or retained_bounds(document)
    for route in document.get('reroutes', []):
        if route.get('internal_transfer'):
            continue
        bad = None
        for a, b in zip(route['points'], route['points'][1:]):
            n = max(1, math.ceil(math.hypot(b['x'] - a['x'], b['z'] - a['z']) / .15))
            for i in range(n + 1):
                f = i / n
                x, z = a['x'] + f * (b['x'] - a['x']), a['z'] + f * (b['z'] - a['z'])
                for zone in document['zones']:
                    if signed_distance_to_zone(x, z, zone) < radius - 1e-4:
                        bad = zone['station_id']; break
                if not bad:
                    for box in auxiliaries:
                        if math.hypot(max(abs(x-box['x'])-box['width']/2,0),max(abs(z-box['z'])-box['depth']/2,0)) < radius-1e-4:
                            bad=box.get('id',box.get('step_id','retained'));break
                if bad:
                    break
            if bad:
                break
        if bad:
            errors.append(issue(route['trip_id'], 'route', f'AGV {radius}m包络进入{bad}'))
    return errors


class PhysicalLayoutStore:
    def __init__(self, directory=None, replay_path=None):
        self.use_active = replay_path is None
        self.directory = Path(directory or os.getenv('SIMPY_PHYSICAL_LAYOUT_STORE', ROOT / 'results/physical_layout/pt6a29-original'))
        self.replay_path = Path(replay_path or os.getenv('SIMPY_REPLAY_PATH', ROOT / 'results/unity_validation/replay.json'))
        self.lock = threading.RLock()
        self.cache = {}
        self.clients = {}

    @property
    def default(self):
        return load_json(DEST / 'layout.json')

    @property
    def current(self):
        path = self.directory / 'current.json'
        return load_json(path) if path.exists() else self.default

    @property
    def replay(self):
        from app.core.unity_active import directory, snapshot
        if self.use_active and (directory()/'active.json').exists():
            return snapshot()
        return load_json(self.replay_path)

    def parse(self, filename, content):
        if len(content) > 2 * 1024 * 1024:
            raise LayoutError([issue('', 'file', '布局文件不能超过2MiB')])
        if filename.lower().endswith('.xlsx'):
            try:
                from openpyxl import load_workbook
                import zipfile
                with zipfile.ZipFile(io.BytesIO(content)) as z:
                    if sum(i.file_size for i in z.infolist()) > 12 * 1024 * 1024:
                        raise ValueError('解压后文件过大')
                book = load_workbook(io.BytesIO(content), read_only=True, data_only=False)
                rows = list(book['设备布局'].iter_rows(values_only=True))
                if list(rows[0]) != SLOT_HEADERS:
                    raise ValueError('请使用原场景29工位专用导出模板')
                source = {s['id']: s for s in self.current['slots']}
                slots = []
                for row in rows[1:]:
                    if not any(v is not None for v in row):
                        continue
                    if any(isinstance(v, str) and v.startswith('=') for v in row):
                        raise ValueError('布局单元格不能为公式')
                    item = dict(zip(SLOT_FIELDS, row))
                    base = copy.deepcopy(source.get(item['id'], {})); base.update(item)
                    slots.append(base)
                result = self.current; result['slots'] = slots
                book.close()
                return result
            except Exception as e:
                raise LayoutError([issue('', 'file', f'Excel读取失败：{e}')]) from e
        if not filename.lower().endswith('.json'):
            raise LayoutError([issue('', 'file', '请选择专用JSON或XLSX布局文件')])
        try:
            return json.loads(content.decode('utf-8-sig'))
        except (ValueError, UnicodeError) as e:
            raise LayoutError([issue('', 'file', f'JSON读取失败：{e}')]) from e

    def validate(self, document):
        errors = []
        base, current = self.default, self.current
        if not isinstance(document, dict):
            raise LayoutError([issue('', 'file', '布局必须是JSON对象')])
        if document.get('profile_id') != PROFILE_ID or document.get('catalog_version') != base['catalog_version']:
            raise LayoutError([issue('', 'profile_id', '请选择原场景29工位专用布局，不能导入旧79设备或独立预览布局')])
        # Topology, process bindings, dimensions, and original recipes are
        # authoritative metadata, not client-editable scheduling interfaces.
        for field in ['zones', 'bindings', 'auxiliary', 'roads', 'gantry', 'obstacles']:
            if not metadata_equal(document.get(field), base[field]) and not metadata_equal(document.get(field), current[field]):
                errors.append(issue('', field, f'{field}为已确认的物理布局定义，不能在设备位置导入中修改'))
        rows = document.get('slots')
        if not isinstance(rows, list):
            raise LayoutError([issue('', 'slots', '缺少设备实例列表')])
        source = {s['id']: s for s in base['slots']}
        zones = {z['station_id']: z for z in base['zones']}
        normalized = copy.deepcopy(base); slots, seen, changes = [], set(), []
        current_slots = {s['id']: s for s in current['slots']}
        live_replay=self.replay if self.use_active else {}
        for raw in rows:
            if not isinstance(raw, dict):
                errors.append(issue('', 'slots', '设备实例必须是对象')); continue
            sid = raw.get('id')
            if sid not in source or sid in seen:
                errors.append(issue(str(sid), 'id', '设备实例未知或重复')); continue
            seen.add(sid); row = copy.deepcopy(source[sid]); valid = True
            for field in ['zone_id', 'model_key', 'representative_id']:
                if raw.get(field) != row[field]:
                    errors.append(issue(sid, field, '所属工位和原模型来源不能替换')); valid = False
            for field in ['x', 'z', 'yaw']:
                try:
                    v = raw[field]
                    if isinstance(v, bool):
                        raise ValueError()
                    value = float(v)
                    if not math.isfinite(value) or abs(value) > 36000:
                        raise ValueError()
                    row[field] = value
                except (ValueError, TypeError, KeyError):
                    errors.append(issue(sid, field, '设备位置和朝向必须为有限数值')); valid = False
            if not valid:
                continue
            delta = math.radians(row['yaw'] - source[sid]['yaw'])
            c, s = math.cos(delta), math.sin(delta)
            for target_x, target_z in [('operator_x', 'operator_z')]:
                dx, dz = source[sid][target_x] - source[sid]['x'], source[sid][target_z] - source[sid]['z']
                row[target_x] = row['x'] + c * dx + s * dz
                row[target_z] = row['z'] - s * dx + c * dz
            row['operator_yaw'] = source[sid]['operator_yaw'] + row['yaw'] - source[sid]['yaw']
            z = zones[row['zone_id']]
            left, top = world(z['u'], z['v']); right, bottom = world(z['u'] + z['width'], z['v'] + z['depth'])
            w, h = row['width'] / 2, row['depth'] / 2
            corners = [(row['x'] + c * x + s * zz, row['z'] - s * x + c * zz)
                       for x, zz in [(-w, -h), (w, -h), (w, h), (-w, h)]]
            if any(not (left + .04 <= x <= right - .04 and bottom + .04 <= zz <= top - .04) for x, zz in corners):
                errors.append(issue(sid, 'position', '设备超出所属工位范围'))
            if not (left + .25 <= row['operator_x'] <= right - .25 and bottom + .25 <= row['operator_z'] <= top - .25):
                errors.append(issue(sid, 'operator', '操作人员位置超出所属工位范围'))
            row['_corners'] = corners
            if any(abs(row[k] - current_slots[sid][k]) > .0001 for k in ['x', 'z', 'yaw']):
                affected = [b['step_id'] for b in base['bindings'] if b['slot_id'] == sid]
                if live_replay.get('schema_version')==4:
                    affected=[n['step_id'] for n in live_replay['process']['nodes'] if (n.get('workstation_id') or n.get('area_id') or n.get('station'))==row['zone_id']]
                changes.append(dict(slot_id=sid, station_id=row['zone_id'], affected_steps=affected,
                                    message=f'{sid}共享设备；影响工序：' + ('、'.join(affected) or '暂无旧流程任务')))
            slots.append(row)
        for sid in sorted(set(source) - seen):
            errors.append(issue(sid, 'id', '缺少设备实例'))
        from app.core.factory_layout import intersects, box
        for i, a in enumerate(slots):
            for b in slots[:i]:
                if a['zone_id'] != b['zone_id']:
                    continue
                if intersects(a['_corners'], b['_corners']):
                    errors.append(issue(a['id'], 'position', f'与{b["id"]}设备重叠'))
                for worker, machine in [(a, b), (b, a)]:
                    if intersects(box(worker['operator_x'], worker['operator_z'], .5, .5), machine['_corners']):
                        errors.append(issue(worker['id'], 'operator', f'操作位置进入{machine["id"]}设备'))
        for row in slots:
            row.pop('_corners', None)
        if errors:
            raise LayoutError(errors)
        normalized['slots'] = slots
        normalized['revision'] = current['revision']
        return dict(valid=True, errors=[], changes=changes, layout=normalized, snapshot=dict(layout=normalized))

    def snapshot(self, layout=None):
        document = copy.deepcopy(layout or self.current)
        from app.core.unity_active import directory,active
        info=active() if self.use_active and (directory()/'active.json').exists() else None
        if info and info['mode']=='pt6a29':
            digest=info['sha256'];sim_id=info['sim_id'];replay=None
        else:
            replay_bytes=self.replay_path.read_bytes();digest=hashlib.sha256(replay_bytes).hexdigest();replay=json.loads(replay_bytes);sim_id=replay['sim_id']
        key = (document['revision'], digest)
        if key not in self.cache:
            replay=replay or self.replay
            if replay.get('schema_version')==4:
                pairs={}
                for trip in replay['transfers']:pairs.setdefault((trip['source_station'],trip['target_station']),trip)
                replay=dict(transfers=list(pairs.values()))
            self.cache = {key: reroutes(document, replay)}
        document['sim_id'] = sim_id
        document['replay_sha256'] = digest
        document['reroutes'] = copy.deepcopy(self.cache[key])
        return document

    def commit(self, document, expected_revision, sim_id):
        with self.lock:
            current = self.current
            if current['revision'] != expected_revision:
                raise RuntimeError('布局版本已变化，请重新预览')
            if self.replay['sim_id'] != sim_id:
                raise RuntimeError('排产已变化，请重新预览')
            normalized = self.validate(document)['layout']
            normalized['revision'] = expected_revision + 1
            atomic_json(self.directory / 'previous.json', current)
            atomic_json(self.directory / f'revision_{expected_revision}.json', current)
            atomic_json(self.directory / 'current.json', normalized)
            self.cache.clear()
            return self.snapshot(normalized)

    def export(self, format='json'):
        if format == 'json':
            return json.dumps(self.snapshot(), ensure_ascii=False, indent=2).encode('utf-8')
        from openpyxl import Workbook
        live=self.replay
        web=live.get('schema_version')==4
        book = Workbook(); sheet = book.active; sheet.title = '设备布局'; sheet.append(SLOT_HEADERS)
        for row in self.current['slots']:
            sheet.append([row[k] for k in SLOT_FIELDS])
        sheet.freeze_panes = 'E2'; sheet.auto_filter.ref = sheet.dimensions
        for col in sheet.columns:
            sheet.column_dimensions[col[0].column_letter].width = 26
        for name, rows, fields in [('旧模型模板绑定' if web else '工序展示安置', self.current['bindings'], ['step_id', 'zone_id', 'slot_id', 'status']),
                                    ('29工位及缓存区', self.current['zones'], ['station_id', 'label', 'u', 'v', 'width', 'depth'])]:
            reference = book.create_sheet(name); reference.append(fields)
            for row in rows:
                reference.append([row[k] for k in fields])
            reference.freeze_panes = 'A2'
            for col in reference.columns:
                reference.column_dimensions[col[0].column_letter].width = 28
        notes = book.create_sheet('填写说明')
        if web:
            sheet=book.create_sheet('当前工步归属');sheet.append(['工步编号','所属物理工位','任务名称'])
            for node in live['process']['nodes']:sheet.append([node['step_id'],node.get('workstation_id') or node.get('area_id') or node.get('station'),node['task_name']])
            sheet.freeze_panes='A2'
            for col in sheet.columns:sheet.column_dimensions[col[0].column_letter].width=30
            notes.append(['当前工步归属来自网页已发布输入快照；旧模型模板绑定只供模型兼容，不是当前工艺映射。'])
        for text in ['仅修改X、Z和水平角度。共享设备的变更影响其全部绑定工序。',
                     '保持设备和操作人员位置在所属29工位内。网页真实工步时间、人员记录和统计结果不修改。',
                     '此表是展示安置表，不是已核定的工艺映射表。']:
            notes.append([text])
        notes.column_dimensions['A'].width = 100
        stream = io.BytesIO(); book.save(stream); return stream.getvalue()

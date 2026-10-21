"""Read process tables without committing a partial flow or changing configuration."""
from __future__ import annotations

import csv
import io
import json
import math
import zipfile
from collections import deque
from pathlib import Path

from pydantic import ValidationError

from app.models.process_model import ProcessDefinition, ProcessNode


REQUIRED_HEADERS = {"step_id", "task_name", "op_type", "std_duration"}
STATION_ALIASES = {"wsation", "wstation", "工位编号"}
MAX_FILE_BYTES = 10 * 1024 * 1024
MAX_XLSX_BYTES = 64 * 1024 * 1024


class ProcessFileError(ValueError):
    def __init__(self, issues):
        self.issues = issues
        self.errors = [
            f"工作表 {i['sheet']}，第 {i['row']} 行，字段 {i['field']}：{i['message']}"
            for i in issues
        ]
        super().__init__("；".join(self.errors))


def issue(sheet, row, field, message):
    return dict(sheet=sheet, row=row, field=field, message=message)


def _text(value):
    return "" if value is None else str(value).strip()


def _header(value):
    value = _text(value).lstrip("\ufeff").casefold()
    return "station" if value in STATION_ALIASES else value


def build_station_catalog(nodes):
    """Use the Unity PT6A catalog for names/order, retaining unknown input IDs."""
    used = {}
    for node in nodes:
        sid = node.workstation_id or node.station
        used.setdefault(sid, dict(station_id=sid, area_id=node.area_id or node.station, label=sid))
    path = Path(__file__).resolve().parents[2] / "data/factory_layout/pt6a29/layout.json"
    try:
        zones = json.loads(path.read_text(encoding="utf-8-sig")).get("zones", [])
    except (OSError, ValueError):
        zones = []
    result = []
    for zone in zones:
        sid = zone.get("station_id")
        if zone.get("kind") == "workstation" and sid in used:
            row = used.pop(sid)
            row["label"] = zone.get("label") or sid
            result.append(row)
    return result + list(used.values())


def _validate_graph(nodes, rows, sheet):
    errors = []
    by_id = {}
    for node, row in zip(nodes, rows):
        if node.step_id in by_id:
            errors.append(issue(sheet, row, "step_id", f"工序编号 {node.step_id} 重复（首次位于第 {by_id[node.step_id][1]} 行）"))
        else:
            by_id[node.step_id] = (node, row)
    if errors:
        return errors
    successors = {sid: [] for sid in by_id}
    indegree = {sid: 0 for sid in by_id}
    for node, row in zip(nodes, rows):
        for predecessor in set(node.get_predecessor_list()):
            if predecessor not in by_id:
                errors.append(issue(sheet, row, "predecessors", f"前置工序 {predecessor} 不存在"))
            else:
                successors[predecessor].append(node.step_id)
                indegree[node.step_id] += 1
    if errors:
        return errors
    ready = deque(sid for sid, degree in indegree.items() if degree == 0)
    while ready:
        sid = ready.popleft()
        for successor in successors[sid]:
            indegree[successor] -= 1
            if not indegree[successor]:
                ready.append(successor)
    remaining = [sid for sid, degree in indegree.items() if degree]
    if remaining:
        # A remaining vertex may be downstream of a cycle. Following an incoming
        # edge within this remainder always reaches a vertex on the actual cycle.
        sid = remaining[0]
        trail = set()
        while sid not in trail:
            trail.add(sid)
            sid = next(pred for pred in by_id[sid][0].get_predecessor_list() if indegree[pred])
        errors.append(issue(sheet, by_id[sid][1], "predecessors", f"流程存在循环依赖，涉及工序 {sid}"))
    return errors


def parse_process_table(headers, rows, *, sheet="CSV", header_row=1, name="未命名流程", xlsx=False):
    columns = [_header(value) for value in headers]
    errors = [issue(sheet, header_row, field, "缺少必需表头") for field in sorted(REQUIRED_HEADERS - set(columns))]
    for field in set(columns) - {""}:
        if columns.count(field) > 1:
            errors.append(issue(sheet, header_row, field, "表头重复或同时使用了同一字段的别名"))
    if errors:
        raise ProcessFileError(errors)
    physical_column = xlsx or any(_text(value).casefold() in STATION_ALIASES for value in headers)
    nodes, source_rows = [], []
    number_defaults = {"std_duration": 0, "time_variance": 0, "work_load_score": 5,
                       "rework_prob": 0, "required_workers": 1, "x": 0, "y": 0}
    for row_number, cells in rows:
        if not any(_text(value) for value in cells):
            continue
        data = {column: cells[index] if index < len(cells) else None for index, column in enumerate(columns) if column}
        row_errors = []
        for field in columns:
            if xlsx and isinstance(data.get(field), str) and data[field].startswith("="):
                row_errors.append(issue(sheet, row_number, field, "不支持公式，请将公式转换为单元格值"))
        for field in ("step_id", "task_name"):
            if not _text(data.get(field)):
                row_errors.append(issue(sheet, row_number, field, "不能为空"))
        if physical_column and not any(_text(data.get(field)) for field in ("station", "area_id", "workstation_id")):
            row_errors.append(issue(sheet, row_number, "station", "工位编号不能为空"))
        if xlsx and not _text(data.get("std_duration")):
            row_errors.append(issue(sheet, row_number, "std_duration", "标准工时不能为空"))
        numbers = {}
        for field, default in number_defaults.items():
            raw = data.get(field)
            try:
                number = default if not _text(raw) else float(raw)
                if isinstance(raw, bool) or not math.isfinite(number):
                    raise ValueError
                if field in ("work_load_score", "required_workers"):
                    if int(number) != number:
                        raise ValueError
                    number = int(number)
                numbers[field] = number
            except (TypeError, ValueError, OverflowError):
                row_errors.append(issue(sheet, row_number, field, "必须为有限数值" if field not in ("work_load_score", "required_workers") else "必须为整数"))
        if row_errors:
            errors.extend(row_errors)
            continue
        station = _text(data.get("station")) or _text(data.get("area_id")) or _text(data.get("workstation_id")) or "ST01"
        bind_physical = physical_column or station.startswith("WS-")
        area_id = _text(data.get("area_id")) or (station if bind_physical else "")
        workstation_id = _text(data.get("workstation_id")) or (station if bind_physical else "")
        try:
            node = ProcessNode(
                step_id=_text(data.get("step_id")), task_name=_text(data.get("task_name")),
                op_type=_text(data.get("op_type")) or "A", predecessors=_text(data.get("predecessors")),
                required_tools=[tool.strip() for tool in _text(data.get("required_tools")).split(";") if tool.strip()],
                station=station, area_id=area_id, workstation_id=workstation_id, **numbers)
        except ValidationError as error:
            errors.extend(issue(sheet, row_number, str(problem["loc"][0]), problem["msg"]) for problem in error.errors())
        else:
            nodes.append(node)
            source_rows.append(row_number)
    errors.extend(_validate_graph(nodes, source_rows, sheet))
    if not nodes and not errors:
        errors.append(issue(sheet, header_row + 1, "step_id", "没有可导入的工艺数据"))
    if errors:
        raise ProcessFileError(errors)
    return ProcessDefinition(name=name, nodes=nodes, station_catalog=build_station_catalog(nodes))


def parse_process_text(content, *, name="未命名流程"):
    try:
        reader = csv.reader(io.StringIO(content))
        header_row, headers = next((reader.line_num, row) for row in reader if any(_text(value) for value in row))
        return parse_process_table(headers, ((reader.line_num, row) for row in reader),
                                   sheet="CSV", header_row=header_row, name=name)
    except (StopIteration, csv.Error) as error:
        raise ProcessFileError([issue("CSV", 1, "file", "CSV 为空或格式无效")]) from error


def parse_process_file(filename, content):
    if len(content) > MAX_FILE_BYTES:
        raise ProcessFileError([issue(filename, 1, "file", "文件不能超过 10 MiB")])
    suffix = Path(filename).suffix.casefold()
    if suffix == ".csv":
        for encoding in ("utf-8-sig", "utf-8", "gbk", "gb2312"):
            try:
                text = content.decode(encoding)
            except UnicodeDecodeError:
                continue
            return parse_process_text(text, name=Path(filename).stem)
        raise ProcessFileError([issue("CSV", 1, "file", "无法识别编码，请使用 UTF-8 或 GBK")])
    if suffix != ".xlsx":
        raise ProcessFileError([issue(filename, 1, "file", "请上传 CSV 或 XLSX 格式文件")])
    book = None
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            if sum(item.file_size for item in archive.infolist()) > MAX_XLSX_BYTES:
                raise ProcessFileError([issue(filename, 1, "file", "XLSX 解压后文件过大")])
        from openpyxl import load_workbook
        book = load_workbook(io.BytesIO(content), read_only=True, data_only=False)
        candidates = []
        first_header = None
        for sheet in book.worksheets:
            rows = enumerate(sheet.iter_rows(values_only=True), 1)
            first = next(((index, values) for index, values in rows if any(_text(value) for value in values)), None)
            if first:
                index, values = first
                first_header = first_header or (sheet.title, index, values)
                if REQUIRED_HEADERS.issubset({_header(value) for value in values}):
                    candidates.append((sheet, index, values, rows))
        if len(candidates) > 1:
            raise ProcessFileError([issue(filename, 1, "sheet", "存在多个工艺工作表，请只保留一个：" + "、".join(row[0].title for row in candidates))])
        if not candidates:
            if first_header:
                title, index, values = first_header
                missing = REQUIRED_HEADERS - {_header(value) for value in values}
                raise ProcessFileError([issue(title, index, field, "缺少必需表头") for field in sorted(missing)])
            raise ProcessFileError([issue(filename, 1, "sheet", "工作簿中没有可导入的工艺工作表")])
        sheet, index, values, rows = candidates[0]
        return parse_process_table(values, rows, sheet=sheet.title, header_row=index, name=Path(filename).stem, xlsx=True)
    except ProcessFileError:
        raise
    except Exception as error:
        raise ProcessFileError([issue(filename, 1, "file", "XLSX 读取失败：" + str(error))]) from error
    finally:
        if book is not None:
            book.close()

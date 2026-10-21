import copy
import io
import json
from pathlib import Path
import pytest
from fastapi import FastAPI
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from app.core.factory_layout import LayoutStore,LayoutError,ASSETS,FIELDS
from app.core.unity_playback import UnityPlaybackClock
from app.api import factory_layout as api

@pytest.fixture
def store(tmp_path):
    source=Path(__file__).resolve().parents[1]/'results/unity_validation/replay.json'
    replay=tmp_path/'replay.json';replay.write_bytes(source.read_bytes())
    return LayoutStore(tmp_path/'layout',replay)

def test_default_roundtrip(store):
    r=store.validate(store.default)
    assert r['valid'] and not r['changes']
    assert len(r['snapshot']['layout']['stations'])==79
    replay=store.replay
    assert len(r['snapshot']['routes'])==len(replay['transfers'])
    assert all(len(t['route'])>=2 for t in r['snapshot']['routes'])

def test_excel_export_read_validation(store):
    data=store.export('xlsx');book=load_workbook(io.BytesIO(data),read_only=False)
    assert book.sheetnames==['工位布局','可选模型','填写说明']
    assert book['工位布局'].max_row==80
    assert len(book['工位布局'].data_validations.dataValidation)>=4
    parsed=store.parse('factory.xlsx',data)
    assert store.validate(parsed)['valid']

@pytest.mark.parametrize('change',[lambda d:d['stations'].pop(),lambda d:d['stations'].append(copy.deepcopy(d['stations'][0])),lambda d:d['stations'][0].update(model_key='missing'),lambda d:d['stations'][0].update(x=float('nan')),lambda d:d['stations'][0].update(scale=0),lambda d:d['stations'][0].update(x=10000),lambda d:d.update(catalog_version='wrong'),lambda d:d['stations'][0].update(show_model='maybe'),lambda d:d['stations'][0].update(animation=None),lambda d:d['stations'][0].update(auto_dock=False,dock_x=0,dock_z=0)])
def test_invalid_never_publish(store,change):
    d=store.default;change(d)
    with pytest.raises(LayoutError):store.validate(d)
    assert not (store.directory/'current.json').exists()

def test_overlap_rejected(store):
    d=store.default;a,b=d['stations'][:2];a['x']=b['x'];a['z']=b['z']
    with pytest.raises(LayoutError,match='重叠'):store.validate(d)

def test_road_rejected(store):
    d=store.default;d['stations'][0].update(x=17.7,z=10)
    with pytest.raises(LayoutError,match='通道'):store.validate(d)

def test_rename_commit_restore_and_replan(store):
    original=store.replay;d=store.default;d['stations'][0]['label']='自定义工位'
    first=store.commit(d,0,original['sim_id'])
    assert first['layout']['revision']==1
    assert store.replay==original
    with pytest.raises(RuntimeError):store.commit(d,0,original['sim_id'])
    replay=store.replay;replay['sim_id']='layout-replan-test';store.replay_path.write_text(json.dumps(replay))
    snap=store.snapshot();assert snap['sim_id']==replay['sim_id']
    assert snap['layout']['stations'][0]['label']=='自定义工位'
    restored=store.commit(store.default,1,replay['sim_id']);assert restored['layout']['revision']==2
    assert (store.directory/'revision_1.json').exists()

def test_arbitrary_model_metadata_ignored(store):
    d=store.default;d['models'][0]['category']='bad';d['stations'][0]['base_width']=100
    result=store.validate(d)['snapshot']['layout']
    assert result['models']==store.default['models']
    assert result['stations'][0]['base_width']==store.default['stations'][0]['base_width']

def test_upload_and_authority_pause_conflicts(store,monkeypatch):
    clock=UnityPlaybackClock(lambda:(store.replay['sim_id'],store.replay['duration']))
    monkeypatch.setattr(api,'store',store);monkeypatch.setattr(api,'clock',clock)
    app=FastAPI();app.include_router(api.router);client=TestClient(app)
    d=store.default;d['stations'][0]['label']='测试工位'
    response=client.post('/layout/validate',files={'file':('layout.json',json.dumps(d).encode(),'application/json')})
    assert response.json()['valid']
    command=dict(client_id='pc',sim_id=store.replay['sim_id'],expected_revision=0,layout=d)
    assert client.put('/layout',json=command).status_code==409
    state=clock.claim('pc');clock.control('pc',state['sim_id'],state['revision'],playing=True)
    assert client.put('/layout',json=command).status_code==409
    state=clock.snapshot();clock.control('pc',state['sim_id'],state['revision'],minute=250,playing=False)
    before=clock.snapshot();assert client.put('/layout',json=command).status_code==200
    after=clock.snapshot();assert before['minute']==after['minute']==250
    assert client.put('/layout',json=command).status_code==409
    assert client.post('/layout/restore',json={k:v for k,v in command.items() if k!='layout'}|dict(expected_revision=1,source='previous')).status_code==200
    assert client.get('/layout/export?format=xlsx').status_code==200
    assert client.get('/layout/previous.json').status_code==200
    ack=client.post('/layout/ack',json=dict(client_id='vr',role='VRViewer',revision=2,sim_id=command['sim_id'])).json()
    assert ack['clients'][0]['revision']==2

def test_formula_and_oversized_upload_rejected(store):
    with pytest.raises(LayoutError):store.parse('huge.json',b'x'*(2*1024*1024+1))
    with pytest.raises(LayoutError):store.parse('layout.xls',b'')

import copy
import io
import json
import struct
from pathlib import Path

import pytest
from fastapi import FastAPI
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.core.factory_layout import LayoutError
from app.core.pt6a29_layout import PhysicalLayoutStore, route_errors
from app.core.unity_playback import UnityPlaybackClock
from app.api import pt6a29_layout as api


@pytest.fixture
def store(tmp_path):
    replay=tmp_path/'replay.json'
    replay.write_bytes((Path(__file__).resolve().parents[1]/'results/unity_validation/replay.json').read_bytes())
    return PhysicalLayoutStore(tmp_path/'physical',replay)


def test_complete_profile_and_route_envelopes(store):
    p=store.snapshot()
    assert store.validate(p)['valid']
    assert len({z['station_id'] for z in p['zones']})==32
    assert sum(z['kind']=='workstation' for z in p['zones'])==29
    assert len(p['bindings'])+len(p['auxiliary'])==79
    assert len({b['step_id'] for b in p['bindings']})==73
    assert {a['step_id'] for a in p['auxiliary']}=={'S001','S002','S003','S004','Q003','Q004'}
    assert not route_errors(p)
    assert any(len(r['points'])>20 for r in p['reroutes'])
    assert all(len(r['points'])>=2 for r in p['reroutes'])


def test_roundtrip_original_workbook_layout_export(store):
    blob=store.export('xlsx');book=load_workbook(io.BytesIO(blob))
    assert book['设备布局'].max_row==81
    assert len(book.sheetnames)==4
    assert store.validate(store.parse('placement.xlsx',blob))['valid']


def f32(value):
    if isinstance(value,dict):return {k:f32(v) for k,v in value.items()}
    if isinstance(value,list):return [f32(v) for v in value]
    if isinstance(value,float):return struct.unpack('f',struct.pack('f',value))[0]
    return value


def test_unity_float_roundtrip_and_canonical_metadata(store):
    result=store.validate(f32(store.default))
    assert result['valid']
    assert result['layout']['zones']==store.default['zones']


@pytest.mark.parametrize('mutation',[
    lambda p:p['zones'][0].update(width=10),
    lambda p:p['bindings'][0].update(step_id='fake'),
    lambda p:p['slots'].pop(),
    lambda p:p['slots'][0].update(x=900),
    lambda p:p['slots'][0].update(yaw=float('nan')),
    lambda p:p['slots'][0].update(model_key='foreign'),
    lambda p:p['slots'][1].update(x=p['slots'][0]['x'],z=p['slots'][0]['z']),
])
def test_invalid_profile_never_overwrites(store,mutation):
    d=store.default;mutation(d)
    with pytest.raises(LayoutError):store.validate(d)
    assert not (store.directory/'current.json').exists()


def test_revision_replan_and_original_replay_unchanged(store):
    before=store.replay_path.read_bytes();p=store.default
    published=store.commit(p,0,store.replay['sim_id'])
    assert published['revision']==1
    assert store.replay_path.read_bytes()==before
    with pytest.raises(RuntimeError):store.commit(p,0,store.replay['sim_id'])
    replay=json.loads(before);replay['sim_id']='review-new-run'
    store.replay_path.write_text(json.dumps(replay),encoding='utf-8')
    assert store.snapshot()['sim_id']=='review-new-run'
    assert store.snapshot()['revision']==1


def test_api_requires_original_clock_authority_pause_and_revision(store,monkeypatch):
    clock=UnityPlaybackClock(lambda:(store.replay['sim_id'],store.replay['duration']))
    monkeypatch.setattr(api,'store',store);monkeypatch.setattr(api,'clock',clock)
    app=FastAPI();app.include_router(api.router);client=TestClient(app)
    p=store.default
    assert client.post('/physical-layout/validate',json=p).json()['valid']
    body=dict(client_id='pc',sim_id=store.replay['sim_id'],expected_revision=0,layout=p)
    assert client.put('/physical-layout',json=body).status_code==409
    state=clock.claim('pc')
    clock.control('pc',body['sim_id'],state['revision'],100,True,10)
    assert client.put('/physical-layout',json=body).status_code==409
    state=clock.snapshot();clock.control('pc',body['sim_id'],state['revision'],100,False,10)
    assert client.put('/physical-layout',json=body).json()['revision']==1
    assert client.put('/physical-layout',json=body).status_code==409
